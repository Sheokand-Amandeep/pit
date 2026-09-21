
"""Create a clean, user-facing Excel workbook and CSV."""

from __future__ import annotations
from pathlib import Path
from typing import Any
import json
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from config import OUTPUT_DIR, EXCEL_FILENAME, CSV_FILENAME

MAIN_COLUMNS = [
    "Disclosure Date", "Company", "Symbol", "Person", "Category", "Action",
    "Shares", "Price Per Share", "Transaction Value", "Transaction Date",
    "Mode", "Exchange", "Shares Before", "Holding % Before",
    "Shares After", "Holding % After", "Submission", "Revision Reason",
]

DETAIL_COLUMNS = [
    "App ID", "Previous App ID", "NSE XML URL", "NSE iXBRL URL", "_Raw XML Fields"
]


def export(rows: list[dict[str, Any]]) -> tuple[Path, Path]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    main_df = pd.DataFrame(rows)
    for col in MAIN_COLUMNS:
        if col not in main_df.columns:
            main_df[col] = None
    main_df = main_df[MAIN_COLUMNS]

    csv_path = OUTPUT_DIR / CSV_FILENAME
    main_df.to_csv(csv_path, index=False, encoding="utf-8-sig")

    excel_path = OUTPUT_DIR / EXCEL_FILENAME
    detail_df = pd.DataFrame(rows)
    for col in DETAIL_COLUMNS:
        if col not in detail_df.columns:
            detail_df[col] = None
    detail_df = detail_df[DETAIL_COLUMNS]

    with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
        main_df.to_excel(writer, sheet_name="PIT Disclosures", index=False)
        detail_df.to_excel(writer, sheet_name="Source Details", index=False)

    _format_workbook(excel_path, len(main_df))
    return excel_path, csv_path


def _format_workbook(path: Path, count: int) -> None:
    wb = load_workbook(path)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 28

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        thin = Side(style="thin", color="D9E1F2")
        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        widths = {}
        for row in ws.iter_rows():
            for cell in row:
                widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value or "")) + 2)

        for col_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width, 12), 42)

    ws = wb["PIT Disclosures"]

    # Human-friendly number formatting.
    header = {cell.value: cell.column for cell in ws[1]}
    for row in range(2, ws.max_row + 1):
        for name in ("Shares", "Shares Before", "Shares After"):
            if name in header:
                ws.cell(row, header[name]).number_format = '#,##0'
        if "Price Per Share" in header:
            ws.cell(row, header["Price Per Share"]).number_format = '₹#,##0.00'
        if "Transaction Value" in header:
            ws.cell(row, header["Transaction Value"]).number_format = '₹#,##0.00'
        for name in ("Holding % Before", "Holding % After"):
            if name in header:
                ws.cell(row, header[name]).number_format = '0.00%'

    if ws.max_row >= 2:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        tab = Table(displayName="PITDisclosures", ref=ref)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tab)

    wb.save(path)
